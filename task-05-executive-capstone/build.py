from pathlib import Path
import csv, subprocess, tempfile, shutil, sys
import pandas as pd
import numpy as np
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image

ROOT = Path(__file__).resolve().parent
DATA = ROOT / "data" / "customer_churn_sample.csv"
OUT = Path(sys.argv[1] if len(sys.argv)>1 else "site/task-05").resolve()
OUT.mkdir(parents=True, exist_ok=True)
df = pd.read_csv(DATA)
required = {"CustomerID","TenureMonths","MonthlyCharges","SupportTickets","ContractType","Churn"}
if required - set(df.columns): raise ValueError("Required fields missing")
if df.CustomerID.isna().any() or df.CustomerID.duplicated().any(): raise ValueError("Customer IDs must be present and unique")
if df[list(required)].isna().any().any(): raise ValueError("Required fields contain missing values")
if not set(df.Churn).issubset({"Yes","No"}): raise ValueError("Churn should contain Yes/No")
n=len(df); ch=df[df.Churn=="Yes"]; rt=df[df.Churn=="No"]
cr=len(ch)/n; exposure=float(ch.MonthlyCharges.sum())
ten_ch=float(ch.TenureMonths.mean()); ten_rt=float(rt.TenureMonths.mean())
tic_ch=float(ch.SupportTickets.mean()); tic_rt=float(rt.SupportTickets.mean())
contract=df.groupby("ContractType").agg(customers=("CustomerID","count"), churned=("Churn",lambda s:int((s=="Yes").sum())))
contract["rate"]=contract.churned/contract.customers
scenarios=[(.10,exposure*.10),(.25,exposure*.25),(.50,exposure*.50)]

# High-contrast, restrained charts.
NAVY="15243B"; INK="1D2D44"; MUTED="61718A"; BG="F4F7FB"; WHITE="FFFFFF"; BLUE="4664D9"; TEAL="178579"; CORAL="C64A60"; GOLD="A9741B"; LINE="E0E7F0"
def hx(v): return RGBColor.from_string(v)
def chart_contract(path):
    d=contract.reindex([x for x in ["Month-to-Month","One Year","Two Year"] if x in contract.index])
    fig,ax=plt.subplots(figsize=(6.3,3.0),dpi=180); fig.patch.set_facecolor("white")
    colors=["#"+(CORAL if x>=.5 else TEAL) for x in d.rate]
    bars=ax.barh(np.arange(len(d)),d.rate*100,color=colors,height=.5)
    ax.set_yticks(np.arange(len(d)),d.index,fontsize=9); ax.set_xlim(0,115)
    ax.set_xlabel("Observed churn rate in sample (%)",fontsize=8)
    ax.grid(axis="x",color="#E0E7F0"); ax.set_axisbelow(True)
    for z in ["top","right","left"]: ax.spines[z].set_visible(False)
    ax.spines["bottom"].set_color("#E0E7F0"); ax.tick_params(axis="y",length=0)
    for b,(_,r) in zip(bars,d.iterrows()): ax.text(b.get_width()+2,b.get_y()+b.get_height()/2,f'{int(r.churned)}/{int(r.customers)} · {r.rate:.0%}',va="center",fontsize=9,fontweight="bold")
    plt.tight_layout(); fig.savefig(path,bbox_inches="tight"); plt.close(fig)
def chart_means(path):
    fig,axs=plt.subplots(1,2,figsize=(6.2,3.0),dpi=180)
    for ax,vals,title in [(axs[0],[ten_ch,ten_rt],"Mean tenure (months)"),(axs[1],[tic_ch,tic_rt],"Mean support tickets / 90d")]:
        bars=ax.bar(["Churned","Not churned"],vals,color=["#"+CORAL,"#"+TEAL],width=.55)
        ax.set_title(title,loc="left",fontsize=10,fontweight="bold"); ax.grid(axis="y",color="#E0E7F0"); ax.set_axisbelow(True)
        for z in ["top","right","left"]:ax.spines[z].set_visible(False)
        ax.spines["bottom"].set_color("#E0E7F0"); ax.tick_params(axis="y",length=0,labelsize=8)
        for b,v in zip(bars,vals):ax.text(b.get_x()+b.get_width()/2,v+max(vals)*.04,f"{v:.1f}",ha="center",fontweight="bold",fontsize=10)
        ax.set_ylim(0,max(vals)*1.25)
    plt.tight_layout(); fig.savefig(path,bbox_inches="tight"); plt.close(fig)
def chart_scenarios(path):
    fig,ax=plt.subplots(figsize=(6.2,2.8),dpi=180)
    vals=[v for _,v in scenarios]; labels=[f"{p:.0%} sensitivity" for p,_ in scenarios]
    bars=ax.barh(labels,vals,color=["#"+BLUE,"#7285E5","#"+TEAL],height=.5)
    ax.invert_yaxis(); ax.grid(axis="x",color="#E0E7F0"); ax.set_axisbelow(True); ax.set_xlabel("INR monthly charge-equivalent (illustrative)",fontsize=8)
    for z in ["top","right","left"]:ax.spines[z].set_visible(False)
    ax.spines["bottom"].set_color("#E0E7F0"); ax.tick_params(axis="y",length=0,labelsize=9)
    for b,v in zip(bars,vals):ax.text(v+max(vals)*.025,b.get_y()+b.get_height()/2,f"₹{v:,.2f}",va="center",fontsize=9,fontweight="bold")
    ax.set_xlim(0,max(vals)*1.45); plt.tight_layout(); fig.savefig(path,bbox_inches="tight"); plt.close(fig)
charts=OUT/"charts"; charts.mkdir(exist_ok=True)
cc=charts/"contract.png"; cm=charts/"means.png"; ci=charts/"scenarios.png"
chart_contract(cc); chart_means(cm); chart_scenarios(ci)

prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5); blank=prs.slide_layouts[6]
def bg(slide,c):
    f=slide.background.fill; f.solid(); f.fore_color.rgb=hx(c)
def box(slide,x,y,w,h,fill,line=None):
    sh=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb=hx(fill); sh.line.fill.background() if not line else None
    if line: sh.line.color.rgb=hx(line); sh.line.width=Pt(1)
    try: sh.adjustments[0]=.1
    except Exception: pass
    return sh
def text(slide,s,x,y,w,h,size=12,color=INK,bold=False,align=PP_ALIGN.LEFT):
    sh=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=sh.text_frame; tf.clear(); tf.word_wrap=True
    tf.margin_left=Inches(.03); tf.margin_right=Inches(.03); tf.margin_top=Inches(.02); tf.margin_bottom=Inches(.02)
    p=tf.paragraphs[0]; p.alignment=align; r=p.add_run(); r.text=str(s); r.font.name="Aptos"; r.font.size=Pt(size); r.font.bold=bold; r.font.color.rgb=hx(color); return sh
def header(slide,num,kicker,title,sub=""):
    text(slide,kicker.upper(),.65,.35,8,.22,9,BLUE,True); text(slide,title,.65,.76,12,.52,24,INK,True)
    if sub:text(slide,sub,.67,1.35,12,.32,10,MUTED)
    text(slide,f"{num:02d} / 08",11.6,7.13,1.1,.16,8,MUTED,True,PP_ALIGN.RIGHT)
def kpi(slide,x,y,w,label,val,foot,col):
    box(slide,x,y,w,1.22,WHITE,LINE); box(slide,x+.14,y+.17,.08,.88,col)
    text(slide,label,x+.32,y+.15,w-.43,.22,8,MUTED,True); text(slide,val,x+.32,y+.43,w-.43,.4,22,INK,True); text(slide,foot,x+.32,y+.91,w-.43,.18,8,MUTED)
# slide 1
s=prs.slides.add_slide(blank); bg(s,NAVY)
text(s,"TASK 05  /  EXECUTIVE ANALYTICS",.8,.72,4,.25,9,"9CAEF9",True)
text(s,"Customer churn,\nwithout the guesswork.",.78,1.42,7.2,1.4,33,WHITE,True)
text(s,"A decision brief for a fictional subscription team",.82,3.15,7,.42,16,"D6E0EF")
box(s,8.6,1.95,3.8,2.2,"233856","455978"); text(s,"OBSERVED SAMPLE CHURN",8.95,2.28,3,.2,9,"D5DFEE",True); text(s,f"{cr:.1%}",8.93,2.7,3,.7,40,WHITE,True); text(s,f"{len(ch)} of {n} records marked churned",8.96,3.53,3,.25,10,"D5DFEE")
text(s,"Recommendation: run a small, measured, service-first retention pilot.",.82,5.23,9,.45,16,WHITE,True); text(s,"Synthetic sample • Descriptive evidence • Uncertainty explicit",.82,6.72,8,.22,10,"BAC8DC")
# slide 2
s=prs.slides.add_slide(blank); bg(s,BG); header(s,2,"01 / Executive decision","Start with service signals, not a prediction","A concise decision for a fictional customer-success lead")
w=2.91
kpi(s,.65,1.9,w,"CUSTOMERS IN SAMPLE",str(n),"Synthetic records",BLUE); kpi(s,3.69,1.9,w,"CHURNED RECORDS",str(len(ch)),f"{cr:.1%} observed rate",CORAL); kpi(s,6.73,1.9,w,"MONTHLY CHARGE EXPOSURE",f"₹{exposure:.2f}","Historical proxy, not a forecast",GOLD); kpi(s,9.77,1.9,w,"AVG TENURE · CHURNED",f"{ten_ch:.1f} mo",f"vs {ten_rt:.1f} mo not churned",TEAL)
box(s,.65,3.5,12.05,2.3,WHITE,LINE); text(s,"THE DECISION",.95,3.78,2,.2,9,BLUE,True); text(s,"Approve a 30-day, human-reviewed support pilot for early-tenure customers with repeated recent tickets.",.95,4.14,10.9,.74,21,INK,True); text(s,"Compare later churn and ticket resolution with a pre-defined baseline before scaling.",.97,5.15,10.8,.3,11,MUTED)
# slide 3
s=prs.slides.add_slide(blank); bg(s,BG); header(s,3,"02 / Five decision-relevant insights","The signal is concentrated in early tenure and support use","All five insights describe this sample and should be validated on more data.")
cards=[(.68,1.9,3.85,1.55,"01 / CHURN","7 of 15","46.7% of records are marked churned.",CORAL), (4.75,1.9,3.85,1.55,"02 / CONTRACT","7 of 7","Month-to-Month marked churned; fixed-term groups 0 of 8.",GOLD), (8.82,1.9,3.85,1.55,"03 / TENURE","7.0 vs 29.1 mo","Mean tenure: churned vs not marked churned.",BLUE), (2.7,3.75,3.85,1.55,"04 / SUPPORT","4.3 vs 1.0","Mean tickets in prior 90 days: churned vs others.",TEAL), (6.77,3.75,3.85,1.55,"05 / EXPOSURE",f"₹{exposure:.2f}/mo","Monthly charges associated with churned records.",GOLD)]
for x,y,w,h,lab,val,body,col in cards:
    box(s,x,y,w,h,WHITE,LINE); box(s,x+.16,y+.18,.08,.99,col); text(s,lab,x+.38,y+.22,w-.55,.2,8,col,True); text(s,val,x+.38,y+.54,w-.55,.42,22,INK,True); text(s,body,x+.38,y+1.08,w-.55,.33,9,MUTED)
text(s,"Use the pattern to select a pilot audience, not to claim what causes churn.",.82,6.2,11.3,.3,12,INK,True)
# slide 4
s=prs.slides.add_slide(blank); bg(s,BG); header(s,4,"03 / What the pattern suggests","Churn appears to cluster around flexible contracts and support use","Observed counts and means are descriptive, not causal.")
box(s,.65,1.8,5.94,4.6,WHITE,LINE); text(s,"Observed churn by contract type",.94,2.05,5,.26,13,INK,True); s.shapes.add_picture(str(cc),Inches(.87),Inches(2.5),width=Inches(5.45),height=Inches(3.2))
box(s,6.78,1.8,5.94,4.6,WHITE,LINE); text(s,"Tenure and support demand by status",7.05,2.05,5.2,.26,13,INK,True); s.shapes.add_picture(str(cm),Inches(7.02),Inches(2.5),width=Inches(5.4),height=Inches(3.2))
# slide 5
s=prs.slides.add_slide(blank); bg(s,BG); header(s,5,"04 / Quantify impact carefully","Use sensitivity ranges for planning, not a forecast","Apply hypothetical percentage reductions to ₹409.93 in monthly charges associated with churned records.")
box(s,.65,1.8,6.65,4.55,WHITE,LINE); text(s,"Illustrative monthly charge-equivalent",.96,2.05,5.9,.25,13,INK,True); s.shapes.add_picture(str(ci),Inches(.9),Inches(2.47),width=Inches(5.98),height=Inches(3.25))
box(s,7.55,1.8,5.12,4.55,NAVY); text(s,"WHAT THESE NUMBERS MEAN",7.9,2.17,4.3,.25,9,"CFD8E7",True); text(s,"10%  →  ₹40.99 / month\n25%  →  ₹102.48 / month\n50%  →  ₹204.97 / month",7.9,2.65,4.3,1.3,18,WHITE,True); text(s,"Assumption: applied arithmetically to the observed exposure proxy.",7.9,4.27,4.15,.7,10,"E0E7F1"); text(s,"Not included: cost, margin, future customer mix, or outcome uncertainty.",7.9,5.25,4.15,.72,10,"F4D9DD")
# slide 6
s=prs.slides.add_slide(blank); bg(s,BG); header(s,6,"05 / Recommended action","Run a bounded 30-day pilot with clear measurement","Learn first; scale only if a larger, fair comparison supports it.")
steps=[(.72,"WEEK 1","DEFINE & REVIEW",BLUE,"Define tenure ≤12 months and ≥3 tickets in prior 90 days. Review unresolved issues and set the baseline."),(4.75,"WEEKS 2–3","CONTACT & RESOLVE",TEAL,"Offer a human-reviewed, service-first check-in. Log contact, response, issue theme, and resolution."),(8.78,"WEEK 4+","COMPARE OUTCOMES",GOLD,"Compare later churn, resolution, feedback, and cost with the pre-defined cohort or baseline.")]
for x,week,title,col,body in steps:
    box(s,x,2.0,3.8,2.65,WHITE,LINE); text(s,week,x+.22,2.23,2,.2,9,col,True); text(s,title,x+.22,2.65,3.3,.27,12,INK,True); text(s,body,x+.22,3.12,3.32,1.1,10,MUTED)
box(s,.72,5.1,11.86,1.0,"E6F5F1"); text(s,"SUCCESS MEASURES",.98,5.34,1.85,.2,8,TEAL,True); text(s,"Subsequent churn • ticket resolution • successful contact • feedback • outreach cost",2.82,5.28,9.1,.4,11,INK,True)
# slide 7
s=prs.slides.add_slide(blank); bg(s,BG); header(s,7,"06 / Evidence boundaries","What this sample can - and cannot - tell us","Transparent uncertainty is part of the recommendation.")
box(s,.68,1.85,5.88,4.5,WHITE,LINE); text(s,"WHAT WE CAN SAY",1.02,2.15,4.7,.25,11,TEAL,True)
for i,line in enumerate(["7 of 15 supplied records are marked churned.","Month-to-Month is 7/7 churned in this sample.","Churned records have shorter mean tenure and more tickets.","Historical monthly charge exposure proxy is ₹409.93."]):
    box(s,1.02,2.7+i*.77,.12,.12,TEAL); text(s,line,1.29,2.61+i*.77,4.65,.48,10,INK)
box(s,6.78,1.85,5.88,4.5,WHITE,LINE); text(s,"WHAT WE CANNOT CLAIM",7.12,2.15,4.7,.25,11,CORAL,True)
for i,line in enumerate(["Representative real-world churn estimate.","That contract or ticket volume causes churn.","Sensitivity figures are guaranteed savings.","Prospective predictive accuracy has been validated."]):
    box(s,7.12,2.7+i*.77,.12,.12,CORAL); text(s,line,7.39,2.61+i*.77,4.65,.48,10,INK)
# slide 8
s=prs.slides.add_slide(blank); bg(s,NAVY); text(s,"DECISION REQUESTED",.82,.63,4.8,.25,10,"9CAEF9",True); text(s,"Approve the pilot.\nMeasure before scaling.",.82,1.25,8.3,1.15,31,WHITE,True); text(s,"Start with recent-tenure customers who have repeated support tickets.",.86,2.75,8.5,.48,15,"DAE3F0")
for i,line in enumerate(["Assign an owner and a fixed follow-up window.","Define a comparison and success measures before contact.","Review churn, issue resolution, feedback, and cost."]):
    box(s,.87,3.55+i*.62,.13,.13,TEAL); text(s,line,1.18,3.49+i*.62,8.2,.35,11,WHITE)
box(s,9.42,1.45,3.0,3.2,"233856","455978"); text(s,"FULL PROJECT",9.72,1.82,2.3,.2,9,"CFD8E7",True); text(s,"Interactive brief\nand supporting files",9.72,2.3,2.3,.72,15,WHITE,True); text(s,"pranav-nc-ai.github.io/customer-churn-data-quality/task-05/",9.72,3.37,2.25,.7,8,"C9D5E6")
text(s,"Five signals. One cautious decision. A measurable next step.",.86,6.38,10,.35,15,"9CAEF9",True)

prs.core_properties.title="Executive Analytics Capstone - Customer Churn"
prs.core_properties.subject="Five insights, impact sensitivities, and a measured retention pilot"
prs.core_properties.author="Data Analytics Internship Project"
pptx=OUT/"executive_analytics_capstone.pptx"; prs.save(pptx)
# Reproducible speaker notes text, one paragraph per slide, passed through eSpeak for an illustrative narrated MP4.
narrations=[
"Today's decision is where to focus a small retention investigation. The supplied sample has only fifteen synthetic customers, so we will not treat it as a forecast. The recommendation is a short service-first pilot focused on early tenure customers showing repeated support activity.",
"Seven of fifteen records are marked churned, an observed rate of forty-six point seven percent. Monthly charges associated with those records total four hundred nine rupees and ninety-three paise. That is a historical exposure proxy, not money we can assume is recoverable. Use it for transparent scenario planning only.",
"Five facts matter. First, seven of fifteen records churned. Second, all seven Month-to-Month records are marked churned, compared with zero of four One Year and zero of four Two Year records. Third, mean tenure is seven versus twenty-nine point one months. Fourth, average recent support tickets are four point three versus one. Fifth, churned records sum to four hundred nine rupees and ninety-three paise in monthly charges.",
"The pattern suggests investigating the early customer experience and recurring support needs. We should ask whether unresolved problems or onboarding friction can be fixed before a renewal decision. These observations do not prove that contract type or tickets cause churn, and the tiny groups make rates unstable.",
"For planning only, applying hypothetical reductions of ten, twenty-five, or fifty percent to the four hundred nine rupees and ninety-three paise exposure proxy gives about forty-one, one hundred two, or two hundred five rupees per month-equivalent. These are sensitivities, not forecasts or promised savings. They exclude cost, margins, future customer mix, and uncertainty.",
"Run a thirty-day pilot. In week one, review ticket themes and define the cohort: tenure of twelve months or less and at least three tickets in the previous ninety days. In weeks two and three, a customer-success owner offers helpful human-reviewed check-ins. Then compare subsequent churn and ticket resolution with a pre-defined baseline.",
"The evidence is limited: fifteen synthetic records, no contact history, no satisfaction measures, no intervention outcomes, and no margin or outreach cost. The analysis is descriptive; it does not establish causality or validate prediction. Any expansion should use a larger cohort, a defined comparison, and a stated follow-up window.",
"The decision request is to approve the small pilot, assign an owner, and agree on success measures before contact begins. Track resolved issues, customer feedback, and later churn; revise or stop if results do not improve. The goal is measured learning, not overconfident prediction."
]
(OUT/"presentation_script.md").write_text("\n\n".join(f"## Slide {i+1}\n\n{t}" for i,t in enumerate(narrations))+"\n",encoding="utf-8")
with (OUT/"impact_scenarios.csv").open("w",newline="",encoding="utf-8") as f:
    w=csv.writer(f); w.writerow(["Assumed reduction","Base monthly charges associated with churn (INR)","Illustrative monthly charge-equivalent (INR)","Interpretation"])
    for pct,val in scenarios:w.writerow([f"{pct:.0%}",f"{exposure:.2f}",f"{val:.2f}","Sensitivity only; not a forecast or guaranteed saving"])
(OUT/"evidence_limitations.md").write_text(f"""# Evidence Limits and Definitions

- Sample size: {n} synthetic records.
- Observed churn rate: {len(ch)} records marked Churn=Yes divided by {n} = {cr:.1%}.
- Monthly charge exposure: sum of MonthlyCharges for records marked churned = ₹{exposure:.2f}. This is a historical proxy, not a forecast.
- Average tenure: {ten_ch:.1f} months for churned records vs {ten_rt:.1f} for records not marked churned.
- Average support tickets in prior 90 days: {tic_ch:.1f} vs {tic_rt:.1f}.
- Hypothetical 10%, 25%, and 50% reductions are arithmetic sensitivities only, not measured treatment effects or guaranteed savings.
- Missing context includes contact history, ticket topics, satisfaction, renewal timing, outcome history, gross margin, outreach cost, and a larger representative sample.
- Recommendation: predefine a comparison group and follow-up period; measure subsequent churn, resolution, feedback, and cost.
""",encoding="utf-8")

# Create the PDF from the same PowerPoint and narrated video from each slide plus its script.
for binary in ["libreoffice","pdftoppm","espeak","ffmpeg"]:
    if not shutil.which(binary): raise RuntimeError(f"Missing required program: {binary}")
subprocess.run(["libreoffice","--headless","--convert-to","pdf","--outdir",str(OUT),str(pptx)],check=True,stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL)
pdf=OUT/"executive_analytics_capstone.pdf"
with tempfile.TemporaryDirectory() as temp:
    t=Path(temp); prefix=t/"slide"
    subprocess.run(["pdftoppm","-png","-r","110",str(pdf),str(prefix)],check=True,stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL)
    pages=sorted(t.glob("slide-*.png"))
    if len(pages)!=8: raise ValueError(f"Expected 8 slides, got {len(pages)}")
    segments=[]
    for i,(page,voice) in enumerate(zip(pages,narrations),1):
        img=Image.open(page).convert("RGB"); img.thumbnail((1280,720),Image.Resampling.LANCZOS)
        canvas=Image.new("RGB",(1280,720),"#F4F7FB"); canvas.paste(img,((1280-img.width)//2,(720-img.height)//2))
        frame=t/f"frame-{i:02}.jpg"; canvas.save(frame,quality=90)
        wav=t/f"audio-{i:02}.wav"; segment=t/f"segment-{i:02}.mp4"
        subprocess.run(["espeak","-v","en-us","-s","158","-p","43","-a","155","-w",str(wav),voice],check=True,stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL)
        subprocess.run(["ffmpeg","-y","-loglevel","error","-loop","1","-framerate","25","-i",str(frame),"-i",str(wav),"-c:v","libx264","-tune","stillimage","-pix_fmt","yuv420p","-c:a","aac","-b:a","128k","-shortest","-vf","scale=1280:720","-movflags","+faststart",str(segment)],check=True)
        segments.append(segment)
    listing=t/"concat.txt"; listing.write_text("\n".join("file '"+str(x)+"'" for x in segments),encoding="utf-8")
    mp4=OUT/"presentation_recording.mp4"
    subprocess.run(["ffmpeg","-y","-loglevel","error","-f","concat","-safe","0","-i",str(listing),"-c","copy","-movflags","+faststart",str(mp4)],check=True)
print(f"Generated {pptx}, {pdf}, {mp4} from {n} source records.")
